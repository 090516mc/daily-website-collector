# -*- coding: utf-8 -*-
"""每日精选网站收集程序（新版逻辑）。

运行流程：
  1. 读取 data/used.json，得到【全部历史记录过的网址】（只记录网址）；
  2. 从网址来源（联网候选源 / 内置兜底池）随机抓取 10 个网址；
  3. 把每次抓到的网址与全部历史记录对比，重复的去掉、补抓，直到凑满 10 个不重复网址；
  4. 调用 DeepSeek（deepseek-v4-flash）模型，用最直白的话介绍每个网址的特点与功能，
     并让模型返回它属于哪一类（学习/工作/娱乐/生活）；
  5. 按四类把介绍生成一篇 Word 文档，保存到 docs/ 目录；
  6. 更新 data/used.json，把这次实际采用的网址追加进历史记录（只记网址）。

需要环境变量：DEEPSEEK_API_KEY（DeepSeek 密钥）。
可选：URL_SOURCE（在线网址源地址）、QQ_MAIL_ADDR + QQ_MAIL_AUTH_CODE（发送邮件）。
"""
import json
import os
import random
import re
import smtplib
import sys
from datetime import datetime
from email.header import Header
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from pathlib import Path

import requests
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from fetcher import fetch_candidates_all

API_URL = "https://api.deepseek.com/chat/completions"
MODEL = "deepseek-v4-flash"
EXPECTED_TOTAL = 10

CATEGORIES = ["学习", "工作", "娱乐", "生活"]

REPO_ROOT = Path(__file__).resolve().parent.parent
DATA_FILE = REPO_ROOT / "data" / "used.json"
DOCS_DIR = REPO_ROOT / "docs"


def get_api_key():
    key = os.environ.get("DEEPSEEK_API_KEY", "").strip()
    if not key:
        sys.exit("错误：没有找到 DEEPSEEK_API_KEY 环境变量，请在脚本或设置里配置 DeepSeek API 密钥。")
    return key


def load_used():
    """历史记录只保存网址。缺省返回 {"urls": [], "updated": ""}。"""
    if DATA_FILE.exists():
        data = json.loads(DATA_FILE.read_text(encoding="utf-8"))
        return {"urls": data.get("urls", []), "updated": data.get("updated", "")}
    return {"urls": [], "updated": ""}


def save_used(used):
    DATA_FILE.parent.mkdir(parents=True, exist_ok=True)
    DATA_FILE.write_text(
        json.dumps(used, ensure_ascii=False, indent=2), encoding="utf-8"
    )


def select_urls(candidates, used_urls, target=EXPECTED_TOTAL):
    """随机抓取 target 个与全部历史记录不重复的网址。

    与记录重复的网址会被丢弃，再从未用过的候选中补抓，直到凑满 target 个；
    若候选源已被历史记录全部覆盖而凑不满，则提示并允许少量重复补齐。
    """
    pool = list(candidates)
    random.shuffle(pool)

    chosen = []
    for url in pool:
        if url in used_urls:
            continue
        chosen.append(url)
        if len(chosen) >= target:
            break

    if len(chosen) < target:
        # 候选源里没有足够“从未用过”的网址了，做保护性补齐
        print(f"警告：候选源可用的新网址不足，目标 {target}，已选出 {len(chosen)} 个，将允许少量重复补齐。")
        for url in pool:
            if url not in chosen:
                chosen.append(url)
            if len(chosen) >= target:
                break
        if len(chosen) < target:
            chosen = pool[:target]
    return chosen


def ask_model(api_key, name, url):
    """调用 DeepSeek 模型，返回介绍文本（含分类）。"""
    prompt = (
        f"请你用最准确、最直白的日常中文（绝对不要使用任何专业术语、技术词汇、英文缩写），"
        f"给普通人介绍下面这个网站。\n"
        f"网站名称：{name}\n网站网址：{url}\n\n"
        f"请严格按照下面的格式输出，每一项各占一行：\n"
        f"【分类】只能填一个，从这四个里选一个：学习、工作、娱乐、生活\n"
        f"【简介】用一两句话说明这个网站是什么、用来干什么\n"
        f"【特点】写出 2 到 3 点它的突出优点\n"
        f"【功能】写出 2 到 3 点它能帮人做什么事\n"
        f"【适合人群】用一句话说明适合什么人使用\n"
        f"不要输出其他任何内容。"
    )
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
    }
    payload = {
        "model": MODEL,
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0.7,
    }
    resp = requests.post(API_URL, json=payload, headers=headers, timeout=60)
    resp.raise_for_status()
    data = resp.json()
    return data["choices"][0]["message"]["content"].strip()


def parse_sections(text):
    """把模型返回的固定格式文本，拆成 {分类, 简介, 特点, 功能, 适合人群}。"""
    sections = {"分类": "生活", "简介": "", "特点": "", "功能": "", "适合人群": ""}
    for key in sections:
        m = re.search(rf"【{key}】(.*?)(?=【|$)", text, re.S)
        if m:
            value = m.group(1).strip()
            if key == "分类":
                value = value.split(",")[0].strip() if value else "生活"
                value = value if value in CATEGORIES else "生活"
            sections[key] = value
    return sections


def build_docx(title, grouped):
    """按四类分组展示。grouped: {分类: [{name,url,text}, ...]}"""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)
    from docx.oxml.ns import qn

    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    head = doc.add_heading(title, level=0)
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER

    intro = doc.add_paragraph(
        "本推荐每天自动生成，随机挑选国内可访问的网站，按学习、工作、娱乐、生活四类归纳，"
        "用最直白的话告诉你每个网站的特点与功能。"
    )
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    def _add_site(name, url, text):
        h = doc.add_heading(name, level=2)
        link = doc.add_paragraph()
        r = link.add_run(url)
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x9C)
        r.underline = True

        if not text:
            doc.add_paragraph("（这次没有拿到介绍内容）")
            return
        sections = parse_sections(text)

        def _add(title_txt, content):
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(title_txt + "：")
            r.bold = True
            p.add_run(content)

        if sections["简介"]:
            _add("简介", sections["简介"])
        if sections["特点"]:
            _add("特点", sections["特点"].replace("\n", "；"))
        if sections["功能"]:
            _add("功能", sections["功能"].replace("\n", "；"))
        if sections["适合人群"]:
            _add("适合人群", sections["适合人群"])
        doc.add_paragraph()

    for category in CATEGORIES:
        items = grouped.get(category)
        if not items:
            continue
        doc.add_heading(category, level=1)
        for it in items:
            _add_site(it["name"], it["url"], it["text"])

    return doc


def send_email(attach_path, title_text):
    """通过 QQ 邮箱 SMTP 把生成的文档作为附件发给收件邮箱。"""
    addr = os.environ.get("QQ_MAIL_ADDR", "").strip()
    code = os.environ.get("QQ_MAIL_AUTH_CODE", "").strip()
    if not addr or not code:
        print("跳过邮件发送：未配置 QQ_MAIL_ADDR 或 QQ_MAIL_AUTH_CODE")
        return

    msg = MIMEMultipart()
    msg["From"] = addr
    msg["To"] = addr
    msg["Subject"] = Header(title_text, "utf-8")
    msg.attach(MIMEText("你好，今天的网站推荐文档已生成，请查收附件。\n", "plain", "utf-8"))

    with open(attach_path, "rb") as f:
        part = MIMEApplication(f.read())
    filename = Path(attach_path).name
    part.add_header("Content-Disposition", "attachment",
                    filename=("utf-8", "", filename))
    msg.attach(part)

    server = smtplib.SMTP_SSL("smtp.qq.com", 465, timeout=60)
    server.login(addr, code)
    server.sendmail(addr, [addr], msg.as_string())
    server.quit()
    print(f"已发送邮件到 {addr}，附件：{filename}")


def main():
    api_key = get_api_key()

    used = load_used()
    used_urls = set(used["urls"])

    candidates = fetch_candidates_all()

    print(f"历史已记录 {len(used_urls)} 个网址，开始抓取 {EXPECTED_TOTAL} 个不重复网址……")
    chosen_urls = select_urls(candidates, used_urls, EXPECTED_TOTAL)
    print(f"本次抓取并去重后，实际采用 {len(chosen_urls)} 个网址。")

    items = []
    for url in chosen_urls:
        info = candidates.get(url, {})
        name = info.get("name") or url
        try:
            text = ask_model(api_key, name, url)
        except Exception as e:
            print(f"  [{name}] 出错：{e}，稍后重试一次……")
            try:
                text = ask_model(api_key, name, url)
            except Exception as e2:
                print(f"  [{name}] 重试仍失败，跳过此网站：{e2}")
                text = ""
        item = {"name": name, "url": url, "text": text}
        item["category"] = parse_sections(text)["分类"] if text else "生活"
        items.append(item)

    grouped = {c: [] for c in CATEGORIES}
    for it in items:
        grouped.setdefault(it["category"], []).append(it)

    today = datetime.now()
    date_label = today.strftime("%Y%m%d")
    title = f"每日精选网站推荐 · {today.strftime('%Y年%m月%d日')}"

    doc = build_docx(title, grouped)
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    out_path = DOCS_DIR / f"每日网站推荐_{date_label}.docx"
    doc.save(str(out_path))
    print(f"已生成：{out_path}")

    try:
        send_email(out_path, title)
    except Exception as e:
        print(f"邮件发送失败（不影响文档已生成）：{e}")

    # 只记录本次实际采用的网址
    used["urls"] = list(dict.fromkeys(list(used["urls"]) + chosen_urls))
    used["updated"] = today.strftime("%Y-%m-%d %H:%M:%S")
    save_used(used)
    print(f"已追加 {len(chosen_urls)} 个网址到历史记录，避免下次重复。")


if __name__ == "__main__":
    main()